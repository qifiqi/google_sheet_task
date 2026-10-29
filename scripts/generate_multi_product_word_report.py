"""按多产品全局预览口径，从 Excel 收益序列生成 Word 回测报告。

Excel 约定：首个工作表、首行为表头；第一列为日期，其余每 3 列一组
``产品名+权重% | index | strat``，index/strat 为该产品的累计收益率。
脚本把每组产品组装为 RPT-M 载荷，直接调用全局预览导出所用的
``strategy_backtest_report_service.generate_word``，不写库、不改任务。

用法（项目根目录执行）：
    .venv\\Scripts\\python scripts\\generate_multi_product_word_report.py A股8只.xlsx
    .venv\\Scripts\\python scripts\\generate_multi_product_word_report.py A股8只.xlsx \\
        --normalize-weights --benchmark 紫金矿业:100 --benchmark 中际旭创:100 \\
        --title "A股组合绩效分析" --output downloads/A股8只报告.docx

要点：
    - 默认按表头权重做绝对权重组合（余额视为现金），与全局预览一致；
      ``--normalize-weights`` 可把权重缩放到合计 100%。
    - ``--benchmark`` 可重复最多 3 次，token 支持产品代码或产品名，
      未选时报告仅含“组合指数”基准列。
    - 未内置映射的产品名会以名称充当代码，平均成交额/净资产自动降级为 "-"。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8")

import openpyxl

DEFAULT_STOCK_CODES = {
    "紫金矿业": "601899.SH",
    "百济神州": "688235.SH",
    "恒立液压": "601100.SH",
    "海光信息": "688041.SH",
    "天齐锂业": "002466.SZ",
    "思源电气": "002028.SZ",
    "深南电路": "002916.SZ",
    "中际旭创": "300308.SZ",
}
HEADER_PATTERN = re.compile(r"^(.*?)\s*([0-9]+(?:\.[0-9]+)?)\s*%$")
MAX_BENCHMARKS = 3
RATIO_QUANT = Decimal("0.0001")


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="从 Excel 多产品累计收益生成全局预览口径的 Word 报告",
    )
    parser.add_argument("xlsx", help="输入 Excel 路径")
    parser.add_argument("--sheet", default="", help="工作表名；缺省取 Sheet1 或首个工作表")
    parser.add_argument("--header-row", type=int, default=1, help="表头所在行（默认 1）")
    parser.add_argument("--output", default="", help="输出 docx 路径；缺省与输入同目录")
    parser.add_argument("--title", default="量化策略回测绩效分析报告", help="报告标题")
    parser.add_argument("--model-version", default="多产品组合", help="报告“模型版本”展示值")
    parser.add_argument("--price-type", default="", help="报告“价格类型”展示值")
    parser.add_argument(
        "--risk-free-rate",
        type=float,
        default=0.0,
        help="年化无风险利率，小数形式（如 0.03 表示 3%%）",
    )
    parser.add_argument(
        "--normalize-weights",
        action="store_true",
        help="把表头权重缩放到合计 100%%；缺省按原值绝对加权（余额视为现金）",
    )
    parser.add_argument(
        "--benchmark",
        action="append",
        default=[],
        metavar="CODE_OR_NAME:RATIO",
        help=f"追加基准列（可重复，最多 {MAX_BENCHMARKS} 条），RATIO 缺省 100",
    )
    parser.add_argument(
        "--no-composite-benchmark",
        action="store_true",
        help="不并列组合指数列（仅在选择自定义基准时生效）",
    )
    parser.add_argument("--market-type", default="cn", help="K 线市场类型（默认 cn）")
    parser.add_argument(
        "--code-map",
        default="",
        help="JSON 文件：{产品名: 股票代码}，覆盖内置映射",
    )
    parser.add_argument("--debug", action="store_true", help="失败时输出完整异常栈")
    return parser.parse_args()


def _parse_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"无法解析日期: {value!r}")


def _load_code_map(path: str) -> dict[str, str]:
    if not path:
        return {}
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError("--code-map 必须是 {产品名: 股票代码} 对象")
    return {str(key).strip(): str(value).strip() for key, value in raw.items()}


def _select_sheet(workbook, sheet_name: str):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"工作表不存在: {sheet_name}")
        return workbook[sheet_name]
    return workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active


def _product_groups(worksheet, header_row: int) -> list[dict]:
    headers = [
        str(worksheet.cell(header_row, column).value or "").strip()
        for column in range(1, worksheet.max_column + 1)
    ]
    groups = []
    for offset, header in enumerate(headers):
        column = offset + 1
        if not header or header.lower() in {"index", "strat"}:
            continue
        if column + 2 > worksheet.max_column:
            continue
        if headers[offset + 1].lower() != "index" or headers[offset + 2].lower() != "strat":
            continue
        match = HEADER_PATTERN.match(header)
        if not match:
            raise ValueError(f"产品表头缺少“名称+权重%”格式: {header!r}（列 {column}）")
        groups.append({
            "name": match.group(1).strip(),
            "ratio": Decimal(match.group(2)),
            "column": column,
        })
    if len(groups) < 2:
        raise ValueError("至少需要 2 个产品列组（产品名+权重%、index、strat）")
    return groups


def _read_returns(worksheet, group: dict, header_row: int) -> list[dict]:
    rows = []
    seen: set[str] = set()
    for row in range(header_row + 1, worksheet.max_row + 1):
        raw_date = worksheet.cell(row, 1).value
        if raw_date is None:
            continue
        day = _parse_date(raw_date)
        index_value = worksheet.cell(row, group["column"] + 1).value
        start_value = worksheet.cell(row, group["column"] + 2).value
        if index_value is None or start_value is None:
            raise ValueError(
                f"{group['name']} 第 {row} 行缺少 index/strat 数据"
            )
        if day in seen:
            raise ValueError(f"{group['name']} 日期重复: {day}")
        seen.add(day)
        rows.append({
            "date": day,
            "index_return": float(index_value),
            "start_return": float(start_value),
        })
    if len(rows) < 2:
        raise ValueError(f"{group['name']} 有效收益行不足 2 天")
    return rows


def _format_ratio(value: Decimal) -> str:
    return f"{format(value.normalize(), 'f')}%"


def _build_products(
    worksheet,
    groups: list[dict],
    header_row: int,
    code_map: dict[str, str],
    normalize: bool,
    market_type: str,
) -> tuple[list[dict], list[str]]:
    ratios = [group["ratio"] for group in groups]
    if normalize:
        total = sum(ratios, Decimal("0"))
        if total <= 0:
            raise ValueError("权重合计必须大于 0，无法归一化")
        ratios = [
            (ratio / total * Decimal("100")).quantize(RATIO_QUANT)
            for ratio in ratios
        ]

    products = []
    unmapped = []
    for group, ratio in zip(groups, ratios):
        code = code_map.get(group["name"], "")
        if not code:
            code = group["name"]
            unmapped.append(group["name"])
        products.append({
            "stock_code": code,
            "product_name": group["name"],
            "ratio": _format_ratio(ratio),
            "market_type": market_type,
            "returns": _read_returns(worksheet, group, header_row),
        })
    return products, unmapped


def _resolve_benchmarks(tokens: list[str], products: list[dict]) -> list[dict]:
    if len(tokens) > MAX_BENCHMARKS:
        raise ValueError(f"基准最多支持 {MAX_BENCHMARKS} 条")
    benchmarks = []
    for token in tokens:
        code, _, ratio_text = token.partition(":")
        code = code.strip()
        ratio_text = ratio_text.strip() or "100"
        percent = Decimal(ratio_text.rstrip("%") or "100").quantize(RATIO_QUANT)
        if percent <= 0 or percent > 100:
            raise ValueError(f"基准比例必须在 (0, 100]: {token}")
        product = next(
            (
                item for item in products
                if code in (item["stock_code"], item["product_name"])
            ),
            None,
        )
        if product is None:
            raise ValueError(f"基准 {code!r} 不在产品列表中")
        benchmarks.append({
            "stock_code": product["stock_code"],
            "ratio": f"{format(percent.normalize(), 'f')}%",
        })
    return benchmarks


def _default_output(xlsx: Path) -> Path:
    return xlsx.with_name(f"{xlsx.stem}_多产品全局预览报告.docx")


def main() -> int:
    args = _parse_args()
    try:
        xlsx = Path(args.xlsx).expanduser().resolve()
        if not xlsx.exists():
            raise ValueError(f"输入文件不存在: {xlsx}")
        if not -1 <= args.risk_free_rate <= 1:
            raise ValueError("--risk-free-rate 必须在 [-1, 1] 区间（0.03 表示 3%）")

        code_map = {**DEFAULT_STOCK_CODES, **_load_code_map(args.code_map)}
        workbook = openpyxl.load_workbook(xlsx, data_only=True)
        worksheet = _select_sheet(workbook, args.sheet)
        groups = _product_groups(worksheet, args.header_row)
        products, unmapped = _build_products(
            worksheet,
            groups,
            args.header_row,
            code_map,
            args.normalize_weights,
            args.market_type,
        )
        benchmarks = _resolve_benchmarks(args.benchmark, products)
        output = Path(args.output).expanduser() if args.output else _default_output(xlsx)
        if output.suffix.lower() != ".docx":
            output = output.with_suffix(".docx")

        first_date = min(item["returns"][0]["date"] for item in products)
        last_date = max(item["returns"][-1]["date"] for item in products)
        print(f"读取 {worksheet.title}: {len(products)} 个产品，{first_date} 至 {last_date}")
        if unmapped:
            print(f"警告: 未映射股票代码的产品将用名称占位: {', '.join(unmapped)}")

        payload = {
            "report_type": "RPT-M",
            "title": args.title,
            "products": products,
            "weighting_mode": "daily_compound",
            "include_composite_benchmark": not args.no_composite_benchmark,
            "metadata": {
                "model_version": args.model_version,
                "price_type": args.price_type,
            },
            "runtime_params": {},
            "filename": output.name,
        }
        if args.risk_free_rate:
            payload["runtime_params"]["risk_free_rate"] = args.risk_free_rate
            payload["metadata"]["risk_free_rate"] = f"{args.risk_free_rate:.2%}"
        if benchmarks:
            payload["index_benchmarks"] = benchmarks

        from app import create_app
        from app.schemas.backtest import StrategyBacktestReportSchema
        from app.services.strategy_backtest_report_service import (
            strategy_backtest_report_service,
        )

        request = StrategyBacktestReportSchema.model_validate(payload)
        app = create_app()
        with app.app_context():
            filename, buffer = strategy_backtest_report_service.generate_word(request)

        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(buffer.getvalue())
        print(f"已生成: {output} ({output.stat().st_size} 字节, 文件名={filename})")
        return 0
    except Exception as exc:
        if args.debug:
            raise
        print(f"生成失败: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

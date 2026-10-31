"""C3 合并导出（/api/exports/tasks/batch-export）Excel 输出行为锁定。"""

import io
import json

from openpyxl import load_workbook

from app.extensions import db
from app.models import Task, TaskResult
from app.services.export_service import export_service


def _task(task_id, name):
    return Task(id=task_id, name=name, task_type="google_sheet", status="completed", config="{}")


def _result(task_id, result_id, step_index=0):
    return TaskResult(
        id=result_id,
        task_id=task_id,
        step_index=step_index,
        parameters=json.dumps({"kline": [{"stock_date": "2025-01-01"}, {"stock_date": "2025-12-31"}]}),
        result=json.dumps({"D2": 0.1, "D3": 0.2}),
        success=True,
    )


def test_batch_export_returns_single_sheet_xlsx(app_factory):
    app = app_factory
    with app.app_context():
        db.session.add(_task("task-a", "任务A"))
        db.session.add(_task("task-b", "任务B"))
        db.session.add(_result("task-a", 1))
        db.session.add(_result("task-b", 2))
        db.session.commit()

        generated = export_service.export_task_results_batch(["task-a", "task-b"])

        assert generated.filename.startswith("C3_合并导出_")
        assert generated.filename.endswith(".xlsx")
        assert generated.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        workbook = load_workbook(io.BytesIO(generated.buffer.getvalue()))
        assert workbook.sheetnames == ["合并导出"]
        # 表头 + 两个任务的各一条结果（与旧 CSV 内容等价）
        rows = list(workbook["合并导出"].iter_rows(values_only=True))
        assert len(rows) == 3


def test_batch_export_keeps_all_rows_of_one_task(app_factory):
    app = app_factory
    with app.app_context():
        db.session.add(_task("task-a", "任务A"))
        db.session.add(_result("task-a", 1))
        db.session.add(_result("task-a", 2))
        db.session.commit()

        generated = export_service.export_task_results_batch(["task-a"])

        workbook = load_workbook(io.BytesIO(generated.buffer.getvalue()))
        rows = list(workbook["合并导出"].iter_rows(values_only=True))
        assert len(rows) == 3  # 表头 + 2 条结果


def test_batch_export_rejects_tasks_without_results(app_factory):
    app = app_factory
    with app.app_context():
        db.session.add(_task("task-empty", "空任务"))
        db.session.commit()

        try:
            export_service.export_task_results_batch(["task-empty"])
        except Exception as exc:  # ValidationError
            assert "均无结果数据" in str(exc)
        else:
            raise AssertionError("空任务应抛出校验错误")

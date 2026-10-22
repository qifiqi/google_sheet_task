from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.exceptions import ValidationError
from app.services.export_workbook_service import (
    C5_EXPORT_COLUMNS,
    build_c7_stock_code_export_archive,
    build_task_export,
)


def _c11_result(result_id, stock_code, random_group, metrics, output_range="AI2:AI26"):
    return {
        "id": result_id,
        "task_id": "task-c11",
        "step_index": result_id,
        "success": True,
        "parameters": {
            "stock_code": stock_code,
            "random_group": random_group,
            "c11_output_range": output_range,
            "kline": [
                {"stock_date": "2024-01-01"},
                {"stock_date": "2024-12-31"},
            ],
        },
        "result": {"sid__模型A": metrics},
    }


C11_METRICS = {
    "AI4": "60.62%",    # off2 return_rate → D2
    "AI5": "9.95%",     # off3 annualized_rate → D3
    "AI6": "-33.78%",   # off4 maxdd → D4
    "AI8": "40.45%",    # off6 index_rate → D5
    "AI9": "7.03%",     # off7 index_annualized_rate → D6
    "AI10": "-32.00%",  # off8 max_index_dd → D7
    "AI12": "2.91%",    # off10 return_beats → D11
    "AI13": "-1.78%",   # off11 dd_beats → D12
    "AI24": "5",        # off22 ref_xm → xm 列
    "flat_result": {
        "start_annual_std_dev": 0.1111,
        "index_annual_std_dev": 0.2222,
        "start_sharpe_ratio": 1.5,
        "index_sharpe_ratio": 2.5,
    },
}


def _c11_task():
    return SimpleNamespace(id="task-c11", name="C11 导出", task_type="google_sheet_C11")


def test_c11_export_maps_result_block_to_c5_columns():
    export = build_task_export(_c11_task(), [_c11_result(7, "600000", "g3", C11_METRICS)])
    sheet = export.workbook.active

    header = [cell.value for cell in sheet[1]]
    assert header[: len(C5_EXPORT_COLUMNS)] == C5_EXPORT_COLUMNS
    # 普通导出列尾：股票代码 + 结果ID
    assert header[len(C5_EXPORT_COLUMNS):] == ["股票代码", "结果ID"]

    row = [cell.value for cell in sheet[2]]
    assert row[:2] == [5, "g3"]
    assert row[2:5] == pytest.approx([0.2017, -0.0178, 0.0292])
    assert row[5:7] == pytest.approx([0.0291, -0.0178])
    assert row[7:13] == pytest.approx([0.6062, 0.0995, -0.3378, 0.4045, 0.0703, -0.32])
    # C11 无单位理论/实际杠杆率收益（D17/D20）→ 留空
    assert row[13:15] == ["", ""]
    assert row[15:19] == pytest.approx([0.1111, 0.2222, 2.5, 1.5])
    assert row[19:] == ["600000", 7]


def test_c11_export_appends_percent_to_bare_numbers():
    metrics = dict(C11_METRICS, AI5="9.95")
    row = [cell.value for cell in build_task_export(_c11_task(), [_c11_result(8, "600000", "g1", metrics)]).workbook.active[2]]
    assert row[8] == pytest.approx(0.0995)


def test_c11_export_honours_custom_output_range():
    # 结果块起点 AK3：偏移量整体换列换行
    metrics = {
        key.replace("AI", "AK"): value
        for key, value in C11_METRICS.items()
        if key.startswith("AI")
    }
    metrics = {key[:-1] + str(int(key[-1]) + 1): value for key, value in metrics.items()}
    metrics["flat_result"] = C11_METRICS["flat_result"]
    results = [_c11_result(9, "600000", "g2", metrics, output_range="AK3:AK27")]

    row = [cell.value for cell in build_task_export(_c11_task(), results).workbook.active[2]]
    assert row[:2] == [5, "g2"]
    assert row[2:5] == pytest.approx([0.2017, -0.0178, 0.0292])


def test_c11_stock_code_export_archive_groups_by_stock():
    results = [
        _c11_result(1, "600000", "g1", C11_METRICS),
        _c11_result(2, "600001", "g2", C11_METRICS),
    ]

    archive = build_c7_stock_code_export_archive(_c11_task(), results)

    assert archive.filename == "C11 导出_按股票代码导出.zip"
    with ZipFile(archive.buffer) as zip_file:
        assert zip_file.namelist() == ["C11 导出_600001.xlsx", "C11 导出_600000.xlsx"]
        workbook = load_workbook(BytesIO(zip_file.read("C11 导出_600000.xlsx")), data_only=True)

    header = [cell.value for cell in workbook.active[1]]
    # 按股票代码导出列尾：仅结果ID
    assert header[len(C5_EXPORT_COLUMNS):] == ["结果ID"]
    assert workbook.active.cell(row=2, column=len(header)).value == 1


def test_stock_code_export_rejects_other_task_types():
    task = SimpleNamespace(id="task-c5", name="C5", task_type="google_sheet_c5")
    with pytest.raises(ValidationError, match="仅支持 C7/C11"):
        build_c7_stock_code_export_archive(task, [_c11_result(1, "600000", "g1", C11_METRICS)])
